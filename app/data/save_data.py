"""XLSX export service using pandas and openpyxl for styled output."""

import logging
import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings

logger = logging.getLogger(settings.PROJECT_NAME)

COLUMN_LABELS = {
    "url": "URL товара",
    "article": "Артикул",
    "name": "Название",
    "price": "Цена (руб.)",
    "images": "Изображения",
    "seller_name": "Продавец",
    "seller_url": "Ссылка на продавца",
    "sizes": "Размеры",
    "stock": "Остаток",
    "rating": "Рейтинг",
    "reviews_count": "Кол-во отзывов",
}

COLUMN_WIDTHS = {
    "url": 45,
    "article": 12,
    "name": 35,
    "price": 14,
    "images": 55,
    "seller_name": 25,
    "seller_url": 40,
    "sizes": 25,
    "stock": 10,
    "rating": 10,
    "reviews_count": 14,
}


def _apply_styles(ws, df: pd.DataFrame) -> None:
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", start_color="2E4057")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = COLUMN_LABELS.get(col_name, col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    wrap_align = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_align
            cell.font = Font(name="Arial", size=10)

    for col_idx, col_name in enumerate(df.columns, start=1):
        width = COLUMN_WIDTHS.get(col_name, 20)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    for row_idx in range(2, len(df) + 2):
        ws.row_dimensions[row_idx].height = 20


def save_xlsx(df: pd.DataFrame, filename: str, sheet_name: str = "Products") -> str:
    os.makedirs(settings.OUTPUT_DIR, exist_ok=True)
    filepath = os.path.join(settings.OUTPUT_DIR, filename)
    logger.info("Saving %d rows to %s", len(df), filepath)
    try:
        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            ws = writer.sheets[sheet_name]
            _apply_styles(ws, df)
        logger.info("File saved successfully: %s", filepath)
    except Exception as exc:
        logger.error("Failed to save file %s: %s", filepath, exc)
        raise
    return filepath


class DataSaver:
    def save_full_catalog(self, df: pd.DataFrame) -> str:
        return save_xlsx(df, settings.FULL_CATALOG_FILENAME, sheet_name="Full Catalog")

    def save_filtered_catalog(self, df: pd.DataFrame) -> str:
        return save_xlsx(df, settings.FILTERED_CATALOG_FILENAME, sheet_name="Filtered")
